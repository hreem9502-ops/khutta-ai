import io
import shutil

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import get_settings
from app.main import create_app
from app.services.plan_spreadsheet import export_plan_workbook
from tests.conftest import clear_dependency_caches

TOKEN = "test-admin-token"


@pytest.fixture
def plans_dir(tmp_path, data_dir):
    target = tmp_path / "plans"
    shutil.copytree(data_dir, target)
    return target


@pytest.fixture
def client(monkeypatch, plans_dir):
    monkeypatch.setenv("ADMIN_TOKEN", TOKEN)
    monkeypatch.setenv("PLANS_DIRECTORY", str(plans_dir))
    monkeypatch.setenv("GEMINI_API_KEY", "")
    monkeypatch.setenv("GEMINI_API_KEY_BACKUP", "")
    clear_dependency_caches()
    yield TestClient(create_app())
    clear_dependency_caches()


def upload(client, data: bytes, *, token: str | None = TOKEN, publish: bool = False):
    headers = {"X-Admin-Token": token} if token else {}
    return client.post(
        f"/api/admin/plans?publish={str(publish).lower()}",
        headers=headers,
        files={"file": ("plan.xlsx", data, "application/octet-stream")},
    )


def renamed_plan(cs_plan, code: str, name_en: str) -> bytes:
    workbook = load_workbook(io.BytesIO(export_plan_workbook(cs_plan)))
    workbook["plan"].cell(2, 2).value = code
    workbook["plan"].cell(4, 2).value = name_en
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_plan_can_be_downloaded_as_a_workbook(client):
    response = client.get("/api/departments/cs/plan.xlsx")

    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert 'filename="cs-study-plan.xlsx"' in response.headers["content-disposition"]


def test_admin_requires_the_token(client, cs_plan):
    data = export_plan_workbook(cs_plan)

    assert upload(client, data, token=None).json()["error"]["code"] == "unauthorized"
    assert upload(client, data, token="wrong").status_code == 401


def test_admin_is_disabled_without_a_configured_token(monkeypatch, cs_plan):
    monkeypatch.setenv("ADMIN_TOKEN", "")
    clear_dependency_caches()
    response = upload(TestClient(create_app()), export_plan_workbook(cs_plan), token="anything")
    clear_dependency_caches()

    assert response.status_code == 403
    assert response.json()["error"]["code"] == "admin_disabled"


def test_check_only_does_not_publish(client, cs_plan, plans_dir):
    response = upload(client, renamed_plan(cs_plan, "se", "Software Engineering"))

    body = response.json()
    assert response.status_code == 200
    assert body["published"] is False
    assert body["summary"]["code"] == "se"
    assert not (plans_dir / "se.json").exists()


def test_publishing_a_valid_plan_makes_it_live_and_persists_it(client, cs_plan, plans_dir):
    response = upload(client, renamed_plan(cs_plan, "se", "Software Engineering"), publish=True)

    assert response.json()["published"] is True
    assert (plans_dir / "se.json").exists()
    departments = {item["code"]: item for item in client.get("/api/departments").json()}
    assert departments["se"]["name"]["en"] == "Software Engineering"


def test_invalid_plan_is_never_published(client, plans_dir):
    response = upload(client, b"not a workbook", publish=True)

    body = response.json()
    assert body["published"] is False
    assert body["issues"][0]["code"] == "invalid_file"
    assert sorted(path.name for path in plans_dir.iterdir()) == ["cs.json", "net.json"]


def test_settings_override_is_isolated():
    clear_dependency_caches()
    assert get_settings().admin_token == ""


def test_read_only_storage_gives_a_clear_error(client, cs_plan, plans_dir, monkeypatch):
    from app.services.plan_store import DirectoryPlanStore

    def read_only(self, plan):
        raise PermissionError("read-only file system")

    monkeypatch.setattr(DirectoryPlanStore, "save", read_only)
    response = upload(client, renamed_plan(cs_plan, "se", "Software Engineering"), publish=True)

    assert response.status_code == 503
    assert response.json()["error"]["code"] == "plan_storage_unavailable"
