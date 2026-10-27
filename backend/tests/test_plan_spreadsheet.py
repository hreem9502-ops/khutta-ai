import io
from collections.abc import Callable

import pytest
from openpyxl import load_workbook

from app.models.admin import PlanIssueCode
from app.services.plan_spreadsheet import export_plan_workbook, read_plan_workbook, summarize_plan


def edited(plan, change: Callable) -> bytes:
    """Export ``plan``, apply ``change(workbook)`` and return the new file."""
    workbook = load_workbook(io.BytesIO(export_plan_workbook(plan)))
    change(workbook)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def course_row(workbook, code: str) -> int:
    sheet = workbook["courses"]
    return next(row[0].row for row in sheet.iter_rows(min_row=2) if row[0].value == code)


def issue_codes(issues) -> set[PlanIssueCode]:
    return {issue.code for issue in issues}


def test_export_then_import_round_trips_the_plan(cs_plan):
    plan, issues = read_plan_workbook(export_plan_workbook(cs_plan))

    assert issues == []
    assert plan == cs_plan


def test_arabic_choice_values_are_accepted(cs_plan):
    def to_arabic(workbook):
        row = course_row(workbook, "CS 510")
        sheet = workbook["courses"]
        sheet.cell(row, 6).value = "اختياري"
        sheet.cell(row, 7).value = "عملي"
        sheet.cell(row, 10).value = "الأول، الثاني"

    plan, issues = read_plan_workbook(edited(cs_plan, to_arabic))

    assert issues == []
    course = plan.by_code["CS 510"]
    assert (course.type, course.category, course.terms) == ("elective", "practical", ["fall", "spring"])


def test_unknown_prerequisite_is_reported_with_its_row(cs_plan):
    def break_prerequisite(workbook):
        workbook["courses"].cell(course_row(workbook, "CS 301"), 8).value = "CS 201, CS 999"

    plan, issues = read_plan_workbook(edited(cs_plan, break_prerequisite))

    assert plan is None
    [issue] = issues
    assert issue.code is PlanIssueCode.UNKNOWN_COURSE
    assert (issue.sheet, issue.field, issue.value) == ("courses", "prerequisites", "CS 999")
    assert issue.row is not None


def test_prerequisite_cycle_is_rejected(cs_plan):
    def make_cycle(workbook):
        workbook["courses"].cell(course_row(workbook, "CS 102"), 8).value = "CS 201"  # CS 201 requires CS 102

    _, issues = read_plan_workbook(edited(cs_plan, make_cycle))

    assert PlanIssueCode.PREREQUISITE_CYCLE in issue_codes(issues)


def test_bad_cells_are_each_reported(cs_plan):
    def corrupt(workbook):
        sheet = workbook["courses"]
        row = course_row(workbook, "CS 101")
        sheet.cell(row, 4).value = "three"
        sheet.cell(row, 6).value = "optional"
        sheet.cell(course_row(workbook, "CS 102"), 2).value = None

    _, issues = read_plan_workbook(edited(cs_plan, corrupt))

    found = {(issue.code, issue.field) for issue in issues}
    assert (PlanIssueCode.INVALID_NUMBER, "credit_hours") in found
    assert (PlanIssueCode.INVALID_CHOICE, "type") in found
    assert (PlanIssueCode.MISSING_VALUE, "name_ar") in found


def test_duplicate_course_codes_are_rejected(cs_plan):
    def duplicate(workbook):
        sheet = workbook["courses"]
        sheet.append([cell.value for cell in sheet[2]])

    _, issues = read_plan_workbook(edited(cs_plan, duplicate))

    assert PlanIssueCode.DUPLICATE_COURSE in issue_codes(issues)


@pytest.mark.parametrize("code", ["../etc", "CS Dept", "x"])
def test_department_code_must_be_a_safe_slug(cs_plan, code):
    def rename(workbook):
        workbook["plan"].cell(2, 2).value = code

    _, issues = read_plan_workbook(edited(cs_plan, rename))

    assert PlanIssueCode.INVALID_DEPARTMENT_CODE in issue_codes(issues)


def test_missing_sheet_and_missing_column(cs_plan):
    def drop_sheet(workbook):
        del workbook["plan"]

    def drop_column(workbook):
        workbook["courses"].cell(1, 5).value = "lvl"

    assert issue_codes(read_plan_workbook(edited(cs_plan, drop_sheet))[1]) == {PlanIssueCode.MISSING_SHEET}
    assert issue_codes(read_plan_workbook(edited(cs_plan, drop_column))[1]) == {PlanIssueCode.MISSING_COLUMN}


@pytest.mark.parametrize("data", [b"", b"not a spreadsheet", b"PK\x03\x04broken"])
def test_non_workbook_upload_is_an_invalid_file(data):
    plan, issues = read_plan_workbook(data)

    assert plan is None
    assert issue_codes(issues) == {PlanIssueCode.INVALID_FILE}


def test_total_hours_are_computed_not_trusted(cs_plan):
    summary = summarize_plan(read_plan_workbook(export_plan_workbook(cs_plan))[0])

    assert summary.total_hours == cs_plan.total_hours
    assert summary.course_count == len(cs_plan.courses)
    assert summary.level_count == 8
