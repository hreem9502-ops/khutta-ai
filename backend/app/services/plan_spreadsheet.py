"""Excel import and export of study plans, so college staff can maintain them.

A workbook has two sheets:

* ``plan``: one ``field | value`` row per department setting (code, names, load rules).
* ``courses``: one row per course with the columns in :data:`COURSE_COLUMNS`.

Choice columns accept English or Arabic values (for example ``elective`` or ``اختياري``),
and lists (prerequisites, terms) may be separated by commas, Arabic commas or semicolons.
"""

import io
import re
import zipfile
from collections.abc import Callable, Iterable
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from app.models.admin import PlanImportSummary, PlanIssue, PlanIssueCode
from app.models.common import normalize_course_code
from app.models.course import CourseCategory, CourseType, StudyPlan, Term
from app.services.plan_validation import COURSES_SHEET, find_plan_issues
from app.services.transcript_analyzer import normalize_token

PLAN_SHEET = "plan"

# (field, help text shown in the exported workbook)
PLAN_FIELDS: list[tuple[str, str]] = [
    ("code", "Department code: lowercase letters, digits or '-' (e.g. cs) / رمز القسم"),
    ("name_ar", "Department name in Arabic / اسم القسم بالعربية"),
    ("name_en", "Department name in English / اسم القسم بالإنجليزية"),
    ("elective_hours_required", "Elective credit hours a student must complete / الساعات الاختيارية المطلوبة"),
    ("gpa_scale", "Maximum GPA, e.g. 4 or 5 / أعلى معدل"),
    ("min_hours", "Minimum registration load / الحد الأدنى للساعات"),
    ("default_max_hours", "Normal maximum load / الحد الأقصى المعتاد"),
    ("high_gpa_threshold", "GPA from which the higher limit applies / المعدل الذي يسمح بحد أعلى"),
    ("high_gpa_max_hours", "Maximum load for high GPA / الحد الأقصى للمعدل المرتفع"),
    ("low_gpa_threshold", "GPA below which the lower limit applies / المعدل الذي يُخفض الحد تحته"),
    ("low_gpa_max_hours", "Maximum load for low GPA / الحد الأقصى للمعدل المنخفض"),
]

COURSE_COLUMNS = [
    "code", "name_ar", "name_en", "credit_hours", "level",
    "type", "category", "prerequisites", "corequisites", "terms",
]
OPTIONAL_COURSE_COLUMNS = {"prerequisites", "corequisites", "terms"}

_DEPARTMENT_CODE = re.compile(r"^[a-z0-9][a-z0-9-]{1,19}$")
_LIST_SEPARATOR = re.compile(r"[,،;\n]+")


def _synonyms(mapping: dict[str, Iterable[str]]) -> dict[str, str]:
    """Map every accepted spelling (normalised) to its canonical value."""
    return {normalize_token(word): canonical for canonical, words in mapping.items() for word in (canonical, *words)}


_TYPES = _synonyms({
    CourseType.REQUIRED: ["إلزامي", "إلزامية", "إجباري", "إجبارية"],
    CourseType.ELECTIVE: ["اختياري", "اختيارية"],
})
_CATEGORIES = _synonyms({
    CourseCategory.THEORETICAL: ["نظري", "نظرية"],
    CourseCategory.PRACTICAL: ["عملي", "عملية"],
    CourseCategory.MIXED: ["مختلط", "مختلطة", "نظري وعملي"],
})
_TERMS = _synonyms({
    Term.FALL: ["الأول", "الفصل الأول", "1"],
    Term.SPRING: ["الثاني", "الفصل الثاني", "2"],
})


# --- Export ---------------------------------------------------------------------


def export_plan_workbook(plan: StudyPlan) -> bytes:
    """Return ``plan`` as an .xlsx file that :func:`read_plan_workbook` can read back."""
    workbook = Workbook()
    settings = workbook.active
    settings.title = PLAN_SHEET
    rules = plan.rules
    values = {
        "code": plan.code,
        "name_ar": plan.name.ar,
        "name_en": plan.name.en,
        "elective_hours_required": plan.elective_hours_required,
        "gpa_scale": rules.gpa_scale,
        "min_hours": rules.min_hours,
        "default_max_hours": rules.default_max_hours,
        "high_gpa_threshold": rules.high_gpa.threshold,
        "high_gpa_max_hours": rules.high_gpa.max_hours,
        "low_gpa_threshold": rules.low_gpa.threshold,
        "low_gpa_max_hours": rules.low_gpa.max_hours,
    }
    settings.append(["field", "value", "notes"])
    for field, note in PLAN_FIELDS:
        settings.append([field, values[field], note])

    courses = workbook.create_sheet(COURSES_SHEET)
    courses.append(COURSE_COLUMNS)
    for course in plan.courses:
        courses.append([
            course.code,
            course.name.ar,
            course.name.en,
            course.credit_hours,
            course.level,
            course.type.value,
            course.category.value,
            ", ".join(course.prerequisites),
            ", ".join(course.corequisites),
            ", ".join(term.value for term in course.terms),
        ])

    for sheet, widths in ((settings, [26, 30, 70]), (courses, [12, 34, 34, 13, 8, 12, 13, 26, 18, 16])):
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for index, width in enumerate(widths):
            sheet.column_dimensions[chr(ord("A") + index)].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- Import ---------------------------------------------------------------------


class _IssueLog:
    """Collects issues while a workbook is read."""

    def __init__(self) -> None:
        self.items: list[PlanIssue] = []

    def add(self, code: PlanIssueCode, sheet: str | None = None, row: int | None = None,
            field: str | None = None, value: Any = None) -> None:
        self.items.append(
            PlanIssue(code=code, sheet=sheet, row=row, field=field, value=None if value is None else str(value))
        )


def read_plan_workbook(data: bytes) -> tuple[StudyPlan | None, list[PlanIssue]]:
    """Parse and validate an uploaded plan. Returns the plan only if there are no issues."""
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError):
        return None, [PlanIssue(code=PlanIssueCode.INVALID_FILE)]

    log = _IssueLog()
    sheets = {name.strip().lower(): workbook[name] for name in workbook.sheetnames}
    for required in (PLAN_SHEET, COURSES_SHEET):
        if required not in sheets:
            log.add(PlanIssueCode.MISSING_SHEET, sheet=required)
    if log.items:
        return None, log.items

    settings = _read_settings(sheets[PLAN_SHEET], log)
    courses, rows = _read_courses(sheets[COURSES_SHEET], log)
    if log.items:
        return None, log.items

    required_hours = sum(course["credit_hours"] for course in courses if course["type"] == CourseType.REQUIRED)
    try:
        plan = StudyPlan.model_validate({
            "code": settings["code"],
            "name": {"ar": settings["name_ar"], "en": settings["name_en"]},
            "total_hours": required_hours + settings["elective_hours_required"],
            "elective_hours_required": settings["elective_hours_required"],
            "rules": {
                "gpa_scale": settings["gpa_scale"],
                "min_hours": settings["min_hours"],
                "default_max_hours": settings["default_max_hours"],
                "high_gpa": {"threshold": settings["high_gpa_threshold"], "max_hours": settings["high_gpa_max_hours"]},
                "low_gpa": {"threshold": settings["low_gpa_threshold"], "max_hours": settings["low_gpa_max_hours"]},
            },
            "courses": courses,
        })
    except ValidationError as error:
        for detail in error.errors():
            log.add(PlanIssueCode.INVALID_RULES, field=".".join(str(part) for part in detail["loc"]))
        return None, log.items

    issues = find_plan_issues(plan, rows)
    return (None, issues) if issues else (plan, [])


def summarize_plan(plan: StudyPlan) -> PlanImportSummary:
    return PlanImportSummary(
        code=plan.code,
        name=plan.name,
        total_hours=plan.total_hours,
        course_count=len(plan.courses),
        level_count=len({course.level for course in plan.courses}),
        elective_count=sum(1 for course in plan.courses if course.type is CourseType.ELECTIVE),
    )


def _read_settings(sheet, log: _IssueLog) -> dict[str, Any]:
    raw: dict[str, tuple[int, Any]] = {}
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if row and row[0] is not None:
            raw[str(row[0]).strip().lower()] = (row_number, row[1] if len(row) > 1 else None)

    def get(field: str, parse: Callable[[Any], Any]) -> Any:
        row_number, value = raw.get(field, (None, None))
        if _blank(value):
            log.add(PlanIssueCode.MISSING_VALUE, PLAN_SHEET, row_number, field)
            return None
        parsed = parse(value)
        if parsed is None:
            log.add(PlanIssueCode.INVALID_NUMBER, PLAN_SHEET, row_number, field, value)
        return parsed

    settings = {
        "code": get("code", lambda value: str(value).strip().lower()),
        "name_ar": get("name_ar", _text),
        "name_en": get("name_en", _text),
    }
    for field in ("elective_hours_required", "min_hours", "default_max_hours", "high_gpa_max_hours", "low_gpa_max_hours"):
        settings[field] = get(field, _int)
    for field in ("gpa_scale", "high_gpa_threshold", "low_gpa_threshold"):
        settings[field] = get(field, _float)

    code = settings["code"]
    if code and not _DEPARTMENT_CODE.match(code):
        log.add(PlanIssueCode.INVALID_DEPARTMENT_CODE, PLAN_SHEET, raw.get("code", (None,))[0], "code", code)
    return settings


def _read_courses(sheet, log: _IssueLog) -> tuple[list[dict[str, Any]], dict[str, int]]:
    rows = sheet.iter_rows(values_only=True)
    header = [str(cell).strip().lower() if cell is not None else "" for cell in next(rows, ())]
    columns = {name: index for index, name in enumerate(header) if name}
    for name in COURSE_COLUMNS:
        if name not in columns and name not in OPTIONAL_COURSE_COLUMNS:
            log.add(PlanIssueCode.MISSING_COLUMN, COURSES_SHEET, 1, name)
    if log.items:
        return [], {}

    courses: list[dict[str, Any]] = []
    row_numbers: dict[str, int] = {}
    for row_number, row in enumerate(rows, start=2):
        if all(_blank(cell) for cell in row):
            continue

        def cell(name: str) -> Any:
            index = columns.get(name)
            return row[index] if index is not None and index < len(row) else None

        course = _read_course_row(cell, row_number, log)
        if course:
            courses.append(course)
            row_numbers.setdefault(course["code"], row_number)
    return courses, row_numbers


def _read_course_row(cell: Callable[[str], Any], row: int, log: _IssueLog) -> dict[str, Any] | None:
    start = len(log.items)

    def required(name: str, parse: Callable[[Any], Any], code: PlanIssueCode) -> Any:
        value = cell(name)
        if _blank(value):
            log.add(PlanIssueCode.MISSING_VALUE, COURSES_SHEET, row, name)
            return None
        parsed = parse(value)
        if parsed is None:
            log.add(code, COURSES_SHEET, row, name, value)
        return parsed

    course = {
        "code": required("code", lambda value: normalize_course_code(str(value)), PlanIssueCode.MISSING_VALUE),
        "name": {
            "ar": required("name_ar", _text, PlanIssueCode.MISSING_VALUE),
            "en": required("name_en", _text, PlanIssueCode.MISSING_VALUE),
        },
        "credit_hours": required("credit_hours", _positive_int, PlanIssueCode.INVALID_NUMBER),
        "level": required("level", _positive_int, PlanIssueCode.INVALID_NUMBER),
        "type": required("type", lambda value: _TYPES.get(normalize_token(str(value))), PlanIssueCode.INVALID_CHOICE),
        "category": required(
            "category", lambda value: _CATEGORIES.get(normalize_token(str(value))), PlanIssueCode.INVALID_CHOICE
        ),
        "prerequisites": [normalize_course_code(code) for code in _split(cell("prerequisites"))],
        "corequisites": [normalize_course_code(code) for code in _split(cell("corequisites"))],
        "terms": _read_terms(cell("terms"), row, log),
    }
    return course if len(log.items) == start else None


def _read_terms(value: Any, row: int, log: _IssueLog) -> list[Term]:
    if _blank(value):
        return [Term.FALL, Term.SPRING]
    terms: list[Term] = []
    for item in _split(value):
        term = _TERMS.get(normalize_token(item))
        if term is None:
            log.add(PlanIssueCode.INVALID_CHOICE, COURSES_SHEET, row, "terms", item)
        elif term not in terms:
            terms.append(term)
    return terms


def _split(value: Any) -> list[str]:
    if _blank(value):
        return []
    return [part.strip() for part in _LIST_SEPARATOR.split(str(value)) if part.strip()]


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text(value: Any) -> str | None:
    text = str(value).strip()
    return text or None


def _int(value: Any) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    return None


def _positive_int(value: Any) -> int | None:
    number = _int(value)
    return number if number is not None and number > 0 else None


def _float(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None
